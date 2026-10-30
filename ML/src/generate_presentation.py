"""
Smart Disease Risk Prediction System
Module: src/generate_presentation.py

Generates a presentation-ready 10-slide PowerPoint presentation (.pptx)
complete with titles, structured bullet points, evaluation tables, and embedded
model evaluation charts.
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


def create_presentation():
    project_root = Path(__file__).resolve().parent.parent
    images_dir = project_root / "static" / "images"
    output_path = project_root / "Smart_Disease_Prediction_Presentation.pptx"

    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    DARK_BLUE = RGBColor(15, 23, 42)     # #0f172a
    PRIMARY_BLUE = RGBColor(37, 99, 235)  # #2563eb
    ACCENT_GREEN = RGBColor(16, 185, 129) # #10b981
    SLATE_GRAY = RGBColor(71, 85, 105)    # #475569
    CARD_BG = RGBColor(241, 245, 249)     # #f1f5f9
    WHITE = RGBColor(255, 255, 255)

    def add_header(slide, title_text, subtitle_text):
        # Header box
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.1))
        tf = title_box.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = DARK_BLUE
        
        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.size = Pt(14)
            p2.font.color.rgb = PRIMARY_BLUE

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = DARK_BLUE
    bg1.line.color.rgb = DARK_BLUE

    t_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(3.8))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "SMART DISEASE RISK PREDICTION SYSTEM"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE

    p2 = tf1.add_paragraph()
    p2.text = "Multi-Model Machine Learning Classification & Health Risk Estimation"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(147, 197, 253)
    p2.space_before = Pt(14)

    p3 = tf1.add_paragraph()
    p3.text = "College ML Project Review & Viva Presentation | 3rd Year B.Tech Computer Science"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(203, 213, 225)
    p3.space_before = Pt(24)

    p4 = tf1.add_paragraph()
    p4.text = "Technologies: Python, scikit-learn, Flask, pandas, NumPy, seaborn"
    p4.font.size = Pt(13)
    p4.font.color.rgb = ACCENT_GREEN
    p4.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 2: Introduction
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "1. Introduction & Motivation", "Background and clinical relevance of early risk identification")
    
    t_box2 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf2 = t_box2.text_frame
    tf2.word_wrap = True

    points2 = [
        ("Global Healthcare Challenge:", "Cardiovascular diseases and metabolic syndromes represent the leading causes of morbidity and mortality worldwide."),
        ("Need for Early Stratification:", "Identifying elevated biomarkers before the onset of acute pathological events enables timely lifestyle and therapeutic interventions."),
        ("Role of Machine Learning:", "Standard clinical scoring systems often assume linear risk relationships. ML captures non-linear interactions across diverse clinical parameters."),
        ("Project Mission:", "Construct an end-to-end, reproducible educational platform that classifies patient risk into LOW RISK or HIGH RISK with calibrated confidence scores."),
        ("Medical Disclaimer:", "Engineered strictly for academic demonstration and research. Not a certified medical diagnostic device.")
    ]
    for title, desc in points2:
        p = tf2.add_paragraph()
        p.text = f"•  {title} "
        p.font.bold = True
        p.font.size = Pt(16)
        p.font.color.rgb = DARK_BLUE
        p.space_before = Pt(14)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = SLATE_GRAY

    # -------------------------------------------------------------
    # SLIDE 3: Problem Statement & Objectives
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "2. Problem Statement & Objectives", "Clear research goals and development targets")

    t_box3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf3 = t_box3.text_frame
    tf3.word_wrap = True

    points3 = [
        ("Problem Statement:", "Design and deploy a reliable, multi-model classification framework capable of processing heterogeneous medical parameters without data leakage."),
        ("Objective 1 - Leakage-Free Preprocessing:", "Construct an 80/20 train/test stratified pipeline utilizing SimpleImputer, StandardScaler, and OneHotEncoder fitted strictly on training data."),
        ("Objective 2 - Multi-Algorithm Benchmarking:", "Train and evaluate 4 diverse baseline classifiers: Logistic Regression, Decision Tree, Random Forest, and Support Vector Machine."),
        ("Objective 3 - Soft Voting Ensemble:", "Develop a meta-classifier aggregating continuous probability vectors to enhance stability and minimize generalization variance."),
        ("Objective 4 - Interactive Full-Stack UI:", "Deploy an accessible Flask web application equipped with instant demo presets, risk gauges, and factor explanations.")
    ]
    for title, desc in points3:
        p = tf3.add_paragraph()
        p.text = f"•  {title} "
        p.font.bold = True
        p.font.size = Pt(16)
        p.font.color.rgb = DARK_BLUE
        p.space_before = Pt(14)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = SLATE_GRAY

    # -------------------------------------------------------------
    # SLIDE 4: Clinical Feature Schema & Dataset
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "3. Health Metrics & Feature Schema", "13 Clinical parameters mapped to binary risk classification")

    t_box4 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.2))
    tf4 = t_box4.text_frame
    tf4.word_wrap = True

    schema_points = [
        ("Demographics:", "Age (20-85 yrs), Gender (Male / Female)."),
        ("Vital Signs:", "Systolic Blood Pressure (mmHg, Normal <120, Hypertensive >=130), Resting Heart Rate (bpm)."),
        ("Laboratory Biomarkers:", "Fasting Glucose (mg/dL, Diabetes threshold >=126), Total Cholesterol (mg/dL, High >=240), BMI (kg/m², Obese >=30)."),
        ("Lifestyle Factors:", "Smoking Status (No / Former / Yes), Physical Activity Level (Low / Moderate / High)."),
        ("Medical History:", "Family History (No / Yes), Diabetes Diagnosis (No / Yes), Hypertension (No / Yes), Prior Heart Disease (No / Yes)."),
        ("Target Variable:", "Binary Target (0 = Low Risk, 1 = High Risk). Centered at median for an exact 50% / 50% balanced distribution across 1,500 records.")
    ]
    for title, desc in schema_points:
        p = tf4.add_paragraph()
        p.text = f"•  {title} "
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = DARK_BLUE
        p.space_before = Pt(12)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = SLATE_GRAY

    # -------------------------------------------------------------
    # SLIDE 5: Preprocessing & Data Leakage Prevention
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "4. Preprocessing Pipeline & Leakage Isolation", "Strict data science standards applied to training and test splits")

    t_box5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.2))
    tf5 = t_box5.text_frame
    tf5.word_wrap = True

    pipe_points = [
        ("Stratified 80/20 Train/Test Partitioning:", "Dataset partitioned into 1,200 training samples and 300 test samples preserving exact 50/50 class balance."),
        ("Strict Prevention of Data Leakage:", "Preprocessing parameters (imputation medians, standard scaling mean/variance) are fitted strictly on X_train. The test set X_test is purely transformed."),
        ("Numerical Feature Pipeline:", "Applied SimpleImputer(strategy='median') followed by StandardScaler() (z-score normalization: z = (x - mu) / sigma)."),
        ("Categorical Feature Pipeline:", "Applied SimpleImputer(strategy='most_frequent') followed by OneHotEncoder(drop='first', handle_unknown='ignore') to avoid multicollinearity."),
        ("Scikit-Learn ColumnTransformer:", "Bundles numerical and categorical pipelines into a unified preprocessor.pkl serialized via joblib for identical production inference.")
    ]
    for title, desc in pipe_points:
        p = tf5.add_paragraph()
        p.text = f"•  {title} "
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = DARK_BLUE
        p.space_before = Pt(12)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = SLATE_GRAY

    # -------------------------------------------------------------
    # SLIDE 6: Machine Learning Models Implemented
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "5. Machine Learning Models & Ensemble Architecture", "Diverse algorithmic approaches combined via Soft Probability Voting")

    t_box6 = slide6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.2))
    tf6 = t_box6.text_frame
    tf6.word_wrap = True

    model_points = [
        ("1. Logistic Regression:", "Linear model modeling log-odds via sigmoid function. Fast benchmark producing well-calibrated baseline probabilities."),
        ("2. Decision Tree Classifier:", "Non-parametric hierarchical rule induction using Gini impurity splits. Regularized at max_depth=5 to avoid overfitting."),
        ("3. Random Forest Classifier:", "Bagging ensemble of 100 decorrelated decision trees with random feature subspaces. Mitigates variance and outputs Gini feature importances."),
        ("4. Support Vector Machine (SVM):", "Identifies maximum-margin hyperplane in high-dimensional space via Radial Basis Function (RBF) kernel. Probabilities fitted via Platt scaling."),
        ("5. Soft Voting Ensemble (Meta-Model):", "Primary classifier. Averages continuous confidence vectors: P_ens = (P_LR + P_DT + P_RF + P_SVM) / 4. Reconciles model disagreement with high stability.")
    ]
    for title, desc in model_points:
        p = tf6.add_paragraph()
        p.text = f"•  {title} "
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = DARK_BLUE
        p.space_before = Pt(12)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = SLATE_GRAY

    # -------------------------------------------------------------
    # SLIDE 7: Results & Model Comparison
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "6. Model Evaluation & Performance Benchmark", "Dynamic evaluation metrics evaluated on 300 held-out test patients")

    # Add Table
    rows = 6
    cols = 6
    left = Inches(0.8)
    top = Inches(1.8)
    width = Inches(6.5)
    height = Inches(2.8)

    table_shape = slide7.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    headers = ["Model", "Accuracy", "Precision", "Recall", "F1 Score", "ROC-AUC"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY_BLUE
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size = Pt(11)
            paragraph.font.bold = True
            paragraph.font.color.rgb = WHITE

    table_data = [
        ["Logistic Reg.", "92.33%", "93.20%", "91.33%", "92.26%", "0.9808"],
        ["Decision Tree", "85.00%", "83.44%", "87.33%", "85.34%", "0.9284"],
        ["Random Forest", "90.00%", "90.00%", "90.00%", "90.00%", "0.9650"],
        ["SVM (RBF)", "91.00%", "90.20%", "92.00%", "91.09%", "0.9780"],
        ["Ensemble (Soft)", "91.33%", "91.33%", "91.33%", "91.33%", "0.9771"]
    ]

    for i, row in enumerate(table_data):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = val
            if i == 4:  # Ensemble row
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(238, 242, 255)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10.5)
                p.font.color.rgb = DARK_BLUE
                if i == 4:
                    p.font.bold = True

    # Embed ROC Chart on right side
    roc_img = images_dir / "roc_curve.png"
    if roc_img.exists():
        slide7.shapes.add_picture(str(roc_img), Inches(7.6), Inches(1.8), width=Inches(5.0))

    # -------------------------------------------------------------
    # SLIDE 8: Feature Importance
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "7. Feature Importance & Clinical Determinants", "Random Forest Gini importance identifying primary risk drivers")

    t_box8 = slide8.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(6.0), Inches(5.0))
    tf8 = t_box8.text_frame
    tf8.word_wrap = True

    fi_points = [
        ("Primary Risk Drivers:", "Gini impurity calculations across 100 decision trees demonstrate that continuous biomarkers and cardiovascular comorbidities hold the highest predictive weight."),
        ("Key Predictors Identified:", "Systolic Blood Pressure, Fasting Glucose, BMI, Age, and Existing Heart Disease contribute predominantly to High Risk categorization."),
        ("Lifestyle Influences:", "Active smoking status and low physical activity accelerate metabolic disease risk."),
        ("Clinical Alignment:", "Matches empirical medical findings where hypertension, diabetes, and hypercholesterolemia act as compounding cardiovascular risk multipliers.")
    ]
    for title, desc in fi_points:
        p = tf8.add_paragraph()
        p.text = f"•  {title} "
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = DARK_BLUE
        p.space_before = Pt(14)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = SLATE_GRAY

    # Embed Feature Importance Chart on right side
    fi_img = images_dir / "feature_importance.png"
    if fi_img.exists():
        slide8.shapes.add_picture(str(fi_img), Inches(7.2), Inches(1.8), width=Inches(5.4))

    # -------------------------------------------------------------
    # SLIDE 9: Web Application & Inference Architecture
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "8. Full-Stack Web Application (Flask)", "Interactive web deployment for clinical decision support simulation")

    t_box9 = slide9.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.2))
    tf9 = t_box9.text_frame
    tf9.word_wrap = True

    web_points = [
        ("Flask MVC Architecture:", "Modular backend controller (app.py) with routes for Assessment, Model Benchmarks, and System Documentation."),
        ("Cached Inference Registry:", "ModelRegistry pattern loads .pkl models into memory upon startup, achieving sub-10ms real-time inference."),
        ("1-Click Demo Presets:", "Pre-programmed 'Normal / Low-Risk' and 'At-Risk' buttons for instant evaluation during college viva presentations."),
        ("Multi-Model Consensus Transparency:", "Displays primary ensemble classification alongside individual model predictions and flags model disagreements."),
        ("Clinical Risk Factor Breakdown:", "Automatically highlights patient biomarkers exceeding reference standards with severity badges.")
    ]
    for title, desc in web_points:
        p = tf9.add_paragraph()
        p.text = f"•  {title} "
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = DARK_BLUE
        p.space_before = Pt(12)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = SLATE_GRAY

    # -------------------------------------------------------------
    # SLIDE 10: Conclusion & Future Scope
    # -------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "9. Conclusion & Future Scope", "Summary of project achievements and prospective extensions")

    t_box10 = slide10.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.2))
    tf10 = t_box10.text_frame
    tf10.word_wrap = True

    concl_points = [
        ("Project Summary:", "Successfully designed, trained, evaluated, and deployed a complete, leak-free disease risk prediction platform achieving 91.33% accuracy and 0.9771 ROC-AUC."),
        ("Ensemble Value:", "Demonstrated that Soft Voting combines the strengths of diverse linear and non-linear classifiers, lowering prediction variance."),
        ("Academic Preparedness:", "Includes automated test suites, dynamic benchmark dashboards, and interactive Jupyter review notebooks."),
        ("Future Scope 1 - Explainable AI:", "Integrate SHAP (Shapley values) and LIME to generate patient-specific waterfall explanation plots."),
        ("Future Scope 2 - Deep Tabular Learning:", "Incorporate TabNet and neural attention architectures for large-scale clinical registries."),
        ("Future Scope 3 - EHR Standards:", "Support HL7/FHIR protocols for direct electronic health record interoperability.")
    ]
    for title, desc in concl_points:
        p = tf10.add_paragraph()
        p.text = f"•  {title} "
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = DARK_BLUE
        p.space_before = Pt(12)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = SLATE_GRAY

    # Save presentation
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")


if __name__ == "__main__":
    create_presentation()
